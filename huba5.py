import sys
import os
from PyQt6.QtWidgets import (QApplication, QMainWindow, QWidget, QVBoxLayout,
                             QHBoxLayout, QPushButton, QLabel, QSpinBox,
                             QDoubleSpinBox, QFileDialog, QTableWidget,
                             QTableWidgetItem, QMessageBox, QHeaderView, QGroupBox,
                             QScrollArea)
from PyQt6.QtCore import Qt
from Bio import SeqIO

# Wykresy i raporty
from matplotlib.backends.backend_qtagg import FigureCanvasQTAgg as FigureCanvas
from matplotlib.figure import Figure
from matplotlib.backends.backend_pdf import PdfPages  # Do obsługi wielostronicowego PDF
import numpy as np

# GUI - STYLE SHEETS
DARK_THEME_STYLE = """
    QMainWindow {
        background-color: #131920;
    }
    QWidget {
        color: #e2e8f0;
        font-family: 'Segoe UI', Tahoma, Geneva, Verdana, sans-serif;
    }
    QGroupBox {
        color: #50E3C2;
        font-weight: bold;
        font-size: 12px;
        border: 1px solid #282f3a;
        border-radius: 6px;
        margin-top: 12px;
        padding-top: 10px;
    }
    QGroupBox::title {
        subcontrol-origin: margin;
        subcontrol-position: top left;
        left: 10px;
        padding: 0 5px;
    }
    QLabel {
        color: #cbd5e0;
        font-size: 12px;
    }
    QSpinBox, QDoubleSpinBox {
        background-color: #0f1319;
        color: #e2e8f0;
        border: 1px solid #282f3a;
        border-radius: 4px;
        padding: 4px 8px;
        min-width: 80px;
    }
    QSpinBox:focus, QDoubleSpinBox:focus {
        border-color: #50E3C2;
    }
    QPushButton {
        background-color: #1a222d;
        color: #cbd5e0;
        border: 1px solid #282f3a;
        border-radius: 4px;
        padding: 6px 12px;
        font-weight: bold;
        font-size: 12px;
    }
    QPushButton:hover {
        background-color: #22323d;
        border-color: #50E3C2;
        color: #50E3C2;
    }
    QPushButton:disabled {
        background-color: #0f1319;
        color: #4a5568;
        border-color: #1a222d;
    }
    QTableWidget {
        background-color: #0f1319;
        color: #e2e8f0;
        gridline-color: #1a222d;
        border: 1px solid #1a222d;
        border-radius: 6px;
    }
    QHeaderView::section {
        background-color: #171d24;
        color: #a0aec0;
        padding: 6px;
        border: 1px solid #1a222d;
        font-weight: bold;
        font-size: 11px;
    }
    QHeaderView::section:hover {
        background-color: #22323d;
        color: #50E3C2;
    }
    QScrollArea {
        border: none;
        background-color: #0f1319;
    }
"""


# FUNKCJE ANALITYCZNE QC
def n_content(seq: str) -> float:
    return seq.count('N') / len(seq) if len(seq) > 0 else 0.0


def gc_content(seq: str) -> float:
    seq_upper = seq.upper()
    g, c = seq_upper.count('G'), seq_upper.count('C')
    return ((g + c) / len(seq_upper)) * 100 if len(seq_upper) > 0 else 0.0


def process_file(file_path: str) -> list[dict]:
    records = []
    ext = file_path.lower()
    if ext.endswith((".fasta", ".fa", ".fna")):
        fmt = "fasta"
    elif ext.endswith((".fastq", ".fq")):
        fmt = "fastq"
    elif ext.endswith((".gbk", ".gb")):
        fmt = "genbank"
    else:
        return records

    try:
        for record in SeqIO.parse(file_path, fmt):
            clean_seq = str(record.seq).upper().replace("-", "")
            records.append({
                "id": record.id,
                "length": len(clean_seq),
                "gc_pct": gc_content(clean_seq),
                "n_pct": n_content(clean_seq),
                "sequence": clean_seq,
                "source": os.path.basename(file_path)
            })
    except Exception as e:
        print(f"Błąd odczytu {file_path}: {e}")
    return records


# GŁÓWNE OKNO APLIKACJI
class PhageQCApp(QMainWindow):
    def __init__(self):
        super().__init__()
        self.setWindowTitle("Phage Host Predictor - Moduł QC - Raportowanie")
        self.setStyleSheet(DARK_THEME_STYLE)
        self.resize(1350, 800)

        self.raw_records = []
        self.accepted_records = []
        self.rejected_records = []

        # Główny kontener pionowy
        central_widget = QWidget()
        self.setCentralWidget(central_widget)
        main_layout = QVBoxLayout(central_widget)
        main_layout.setContentsMargins(15, 15, 15, 15)
        main_layout.setSpacing(10)

        # ================= SEKCJA GÓRNA: PARAMETRY I AKCJE =================
        top_bar_layout = QHBoxLayout()

        # Grupa: Parametry Filtrowania
        filter_group = QGroupBox("Parametry Filtrowania")
        filter_layout = QHBoxLayout(filter_group)
        filter_layout.setSpacing(10)

        filter_layout.addWidget(QLabel("Min. długość (bp):"))
        self.min_len_spin = QSpinBox()
        self.min_len_spin.setRange(1, 1000000)
        self.min_len_spin.setValue(5000)
        self.min_len_spin.setSingleStep(500)
        filter_layout.addWidget(self.min_len_spin)

        filter_layout.addWidget(QLabel("Maks. 'N' (%):"))
        self.max_n_spin = QDoubleSpinBox()
        self.max_n_spin.setRange(0.0, 100.0)
        self.max_n_spin.setValue(5.0)
        filter_layout.addWidget(self.max_n_spin)

        top_bar_layout.addWidget(filter_group, stretch=2)

        # Grupa: Akcje i Raporty
        actions_group = QGroupBox("Akcje i Raporty")
        actions_layout = QHBoxLayout(actions_group)
        actions_layout.setSpacing(8)

        self.btn_load = QPushButton("Wgraj pliki...")
        self.btn_load.clicked.connect(self.load_files)
        actions_layout.addWidget(self.btn_load)

        self.btn_run = QPushButton("Analizuj (Run QC)")
        self.btn_run.setEnabled(False)
        self.btn_run.setStyleSheet("color: #50E3C2; border-color: #50E3C2;")
        self.btn_run.clicked.connect(self.run_qc)
        actions_layout.addWidget(self.btn_run)

        self.btn_export = QPushButton("Eksportuj FASTA")
        self.btn_export.setEnabled(False)
        actions_layout.addWidget(self.btn_export)
        self.btn_export.clicked.connect(self.export_fasta)

        self.btn_excel = QPushButton("Raport Excel (.xlsx)")
        self.btn_excel.setEnabled(False)
        actions_layout.addWidget(self.btn_excel)
        self.btn_excel.clicked.connect(self.export_excel)

        self.btn_pdf = QPushButton("Raport PDF (.pdf)")
        self.btn_pdf.setEnabled(False)
        actions_layout.addWidget(self.btn_pdf)
        self.btn_pdf.clicked.connect(self.export_pdf)

        top_bar_layout.addWidget(actions_group, stretch=3)
        main_layout.addLayout(top_bar_layout)

        self.lbl_summary = QLabel("Wszystkie: 0 | Zaakceptowane: 0 | Odrzucone: 0")
        self.lbl_summary.setStyleSheet("font-weight: bold; color: #a0aec0; padding-left: 2px;")
        main_layout.addWidget(self.lbl_summary)

        content_layout = QHBoxLayout()

        self.table = QTableWidget()
        self.table.setColumnCount(7)
        self.table.setHorizontalHeaderLabels([
            "ID Sekwencji", "Długość", "GC (%)", "N (%)", "Quality Score", "Status", "Powód / Źródło"
        ])
        self.table.horizontalHeader().setSectionResizeMode(QHeaderView.ResizeMode.Stretch)
        self.table.setSortingEnabled(True)

        content_layout.addWidget(self.table, stretch=3)

        charts_sidebar = QVBoxLayout()
        charts_sidebar.setSpacing(10)

        # Wykres kołowy podsumowania
        self.pie_fig = Figure(figsize=(4, 3), dpi=100, facecolor='#0f1319')
        self.pie_canvas = FigureCanvas(self.pie_fig)
        self.pie_canvas.setStyleSheet("border: 1px solid #1a222d; border-radius: 4px;")
        charts_sidebar.addWidget(self.pie_canvas, stretch=1)

        # Wykres profilu (Heatmapa) umieszczony w QScrollArea
        self.heatmap_fig = Figure(figsize=(4, 4), dpi=100, facecolor='#0f1319')
        self.heatmap_canvas = FigureCanvas(self.heatmap_fig)

        self.heatmap_scroll = QScrollArea()
        self.heatmap_scroll.setWidgetResizable(True)
        self.heatmap_scroll.setWidget(self.heatmap_canvas)
        self.heatmap_scroll.setStyleSheet("border: 1px solid #1a222d; border-radius: 4px;")
        charts_sidebar.addWidget(self.heatmap_scroll, stretch=1)

        content_layout.addLayout(charts_sidebar, stretch=1)
        main_layout.addLayout(content_layout)

        self.update_pie_chart(0, 0)
        self.update_heatmap([])

    # LOGIKA PROGRAMU
    def load_files(self):
        files, _ = QFileDialog.getOpenFileNames(
            self, "Wybierz pliki sekwencji", "",
            "Pliki sekwencji (*.fasta *.fa *.fna *.fastq *.fq *.gbk);;Wszystkie pliki (*)"
        )
        if not files: return

        self.raw_records = []
        for file_path in files:
            self.raw_records.extend(process_file(file_path))

        if not self.raw_records:
            QMessageBox.warning(self, "Błąd", "Nie znaleziono poprawnych sekwencji.")
            self.btn_run.setEnabled(False)
            return

        self.btn_run.setEnabled(True)
        self.lbl_summary.setText(f"Wczytano z plików: {len(self.raw_records)} sekwencji. Kliknij 'Analizuj (Run QC)'.")

    def run_qc(self):
        if not self.raw_records: return

        self.table.setSortingEnabled(False)

        self.accepted_records = []
        self.rejected_records = []
        allowed_chars = set("ACGTN")

        min_len = self.min_len_spin.value()
        max_n_pct = self.max_n_spin.value() / 100.0

        for r in self.raw_records:
            n_factor = max(0.0, 100.0 - (r["n_pct"] * 100 * 10))
            gc = r["gc_pct"]
            gc_factor = 100.0 if 35 <= gc <= 65 else max(0.0, 100.0 - (min(abs(gc - 35), abs(gc - 65)) * 4))
            len_factor = max(0.0, min(100.0, (r["length"] / 25000) * 100))
            r["quality_score"] = round((n_factor * 0.50) + (gc_factor * 0.25) + (len_factor * 0.25), 1)

            invalid_chars = set(r["sequence"]) - allowed_chars
            if invalid_chars:
                r["reject_reason"] = "Niewłaściwe znaki alfabetu"
                self.rejected_records.append(r)
            elif r["length"] < min_len:
                r["reject_reason"] = f"Za krótka (< {min_len} bp)"
                self.rejected_records.append(r)
            elif r["n_pct"] > max_n_pct:
                r["reject_reason"] = f"Zbyt wysokie N (> {max_n_pct * 100:.1f}%)"
                self.rejected_records.append(r)
            else:
                self.accepted_records.append(r)

        self.lbl_summary.setText(
            f"Wszystkie: {len(self.raw_records)} | "
            f"Zaakceptowane: {len(self.accepted_records)} | "
            f"Odrzucone: {len(self.rejected_records)}"
        )

        self.table.setRowCount(len(self.raw_records))
        row = 0
        for r in self.accepted_records:
            self._add_table_row(row, r, "Zaakceptowano", r["source"], Qt.GlobalColor.green)
            row += 1
        for r in self.rejected_records:
            self._add_table_row(row, r, "Odrzucono", r["reject_reason"], Qt.GlobalColor.red)
            row += 1
        self.table.setSortingEnabled(True)

        self.update_pie_chart(len(self.accepted_records), len(self.rejected_records))
        self.update_heatmap(self.raw_records)

        self.btn_export.setEnabled(len(self.accepted_records) > 0)
        self.btn_excel.setEnabled(len(self.raw_records) > 0)
        self.btn_pdf.setEnabled(len(self.raw_records) > 0)

    def _add_table_row(self, row, record, status_text, reason_text, color):
        id_item = QTableWidgetItem(record["id"])

        len_item = QTableWidgetItem()
        len_item.setData(Qt.ItemDataRole.DisplayRole, int(record["length"]))

        gc_item = QTableWidgetItem()
        gc_item.setData(Qt.ItemDataRole.DisplayRole, float(f"{record['gc_pct']:.2f}"))

        n_item = QTableWidgetItem()
        n_item.setData(Qt.ItemDataRole.DisplayRole, float(f"{record['n_pct'] * 100:.2f}"))

        qs_item = QTableWidgetItem()
        qs_item.setData(Qt.ItemDataRole.DisplayRole, float(f"{record['quality_score']:.1f}"))

        status_item = QTableWidgetItem(status_text)
        status_item.setForeground(color)
        reason_item = QTableWidgetItem(reason_text)

        self.table.setItem(row, 0, id_item)
        self.table.setItem(row, 1, len_item)
        self.table.setItem(row, 2, gc_item)
        self.table.setItem(row, 3, n_item)
        self.table.setItem(row, 4, qs_item)
        self.table.setItem(row, 5, status_item)
        self.table.setItem(row, 6, reason_item)

    # GENEROWANIE WYKRESÓW
    def update_pie_chart(self, accepted, rejected):
        self.pie_fig.clear()
        ax = self.pie_fig.add_subplot(111)
        ax.set_facecolor('#0f1319')

        if accepted == 0 and rejected == 0:
            ax.text(0.5, 0.5, "Podsumowanie Filtracji\n(Brak danych)", ha='center', va='center', color='#a0aec0')
            ax.axis('off')
        else:
            labels = [l for l, s in zip(['Zaakceptowane', 'Odrzucone'], [accepted, rejected]) if s > 0]
            sizes = [s for s in [accepted, rejected] if s > 0]
            colors = [c for c, s in zip(['#2ecc71', '#e74c3c'], [accepted, rejected]) if s > 0]

            ax.pie(sizes, labels=labels, colors=colors, autopct='%1.1f%%', startangle=140,
                   textprops={'fontsize': 9, 'color': 'white', 'weight': 'bold'})
            ax.set_title("Podsumowanie Filtracji", color='#cbd5e0', fontsize=11, weight='bold')
            ax.axis('equal')

        self.pie_fig.tight_layout()
        self.pie_canvas.draw()

    def update_heatmap(self, records):
        self.heatmap_fig.clear()
        ax = self.heatmap_fig.add_subplot(111)
        ax.set_facecolor('#0f1319')

        if not records:
            ax.text(0.5, 0.5, "Profil Parametrów QC\n(Uruchom analizę)", ha='center', va='center', color='#a0aec0')
            ax.axis('off')
            self.heatmap_canvas.draw()
            return

        labels = [r["id"][:10] for r in records]
        gc_vals = [r["gc_pct"] for r in records]
        n_vals = [r["n_pct"] * 100 for r in records]
        qs_vals = [r["quality_score"] for r in records]

        # Dynamiczne skalowanie wysokości wykresu pod scrollbar
        dynamic_height = max(4.0, len(records) * 0.3)
        self.heatmap_fig.set_size_inches(4, dynamic_height)

        data_matrix = np.array([gc_vals, n_vals, qs_vals]).T

        color_matrix = np.zeros_like(data_matrix, dtype=float)
        for j in range(3):
            cmax = data_matrix[:, j].max()
            color_matrix[:, j] = data_matrix[:, j] / cmax if cmax > 0 else 0

        ax.imshow(color_matrix, cmap='YlGnBu', aspect='auto')
        ax.set_yticks(np.arange(len(labels)))
        ax.set_yticklabels(labels, fontsize=8, color='white')
        ax.set_xticks(np.arange(3))

        # POPRAWKA: Przeniesienie nazw kolumn na górę i dodanie marginesu do tytułu
        ax.xaxis.tick_top()
        ax.set_xticklabels(['% GC', '% N', 'Quality'], fontsize=9, color='white')
        ax.set_title("Profil Parametrów QC", color='#cbd5e0', fontsize=11, weight='bold', pad=25)

        for i in range(len(labels)):
            for j in range(3):
                val = data_matrix[i, j]
                ax.text(j, i, f"{val:.1f}", ha='center', va='center',
                        color='white' if color_matrix[i, j] > 0.5 else 'black', fontsize=8, weight='bold')

        self.heatmap_fig.tight_layout()
        self.heatmap_canvas.draw()
        self.heatmap_canvas.setMinimumHeight(int(dynamic_height * 100))

    # EKSPORTY I RAPORTY
    def export_excel(self):
        all_data = self.accepted_records + self.rejected_records
        if not all_data: return
        save_path, _ = QFileDialog.getSaveFileName(self, "Zapisz Excel", "Raport_QC.xlsx", "Skoroszyt Excel (*.xlsx)")
        if not save_path: return
        try:
            import openpyxl
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Raport QC"
            ws.append(["ID", "Długość", "GC (%)", "N (%)", "Quality Score", "Status", "Powód"])
            for r in all_data:
                status = "Zaakceptowano" if r in self.accepted_records else "Odrzucono"
                ws.append([r["id"], r["length"], round(r["gc_pct"], 2), round(r["n_pct"] * 100, 2), r["quality_score"],
                           status, r.get("reject_reason", r["source"])])
            wb.save(save_path)
            QMessageBox.information(self, "Sukces", "Pomyślnie wyeksportowano arkusz Excel!")
        except Exception as e:
            QMessageBox.critical(self, "Błąd", f"Nie udało się zapisać pliku: {e}")

    def export_pdf(self):
        all_data = self.accepted_records + self.rejected_records
        if not all_data:
            QMessageBox.warning(self, "Błąd", "Brak danych do wygenerowania raportu.")
            return

        # Okno wyboru ścieżki zapisu PDF
        save_path, _ = QFileDialog.getSaveFileName(self, "Zapisz Raport PDF", "Raport_QC_Pelny.pdf",
                                                   "Pliki PDF (*.pdf)")
        if not save_path:
            return

        try:
            with PdfPages(save_path) as pdf:
                # 1. STRONA 1: Wykres kołowy podsumowania filtracji
                pdf.savefig(self.pie_fig)

                # 2. STRONA 2: Cały wykres profilu parametrów (Heatmapa)
                pdf.savefig(self.heatmap_fig)

                # 3. STRONY KOLEJNE: Tabele z wynikami sekwencji (format Landscape)
                rows_per_page = 22
                col_labels = ["ID Sekwencji", "Długość (bp)", "GC (%)", "N (%)", "Quality Score", "Status",
                              "Powód / Źródło"]
                col_widths = [0.18, 0.10, 0.09, 0.09, 0.11, 0.14, 0.29]

                for i in range(0, len(all_data), rows_per_page):
                    chunk = all_data[i:i + rows_per_page]

                    fig_table = Figure(figsize=(11, 8.5), dpi=100)
                    ax_table = fig_table.add_subplot(111)
                    ax_table.axis('tight')
                    ax_table.axis('off')

                    cell_text = []
                    for r in chunk:
                        status = "Zaakceptowano" if r in self.accepted_records else "Odrzucono"
                        reason = r.get("reject_reason", r["source"])
                        cell_text.append([
                            str(r["id"]),
                            str(r["length"]),
                            f"{r['gc_pct']:.2f}",
                            f"{r['n_pct'] * 100:.2f}",
                            f"{r['quality_score']:.1f}",
                            status,
                            str(reason)
                        ])

                    table_obj = ax_table.table(
                        cellText=cell_text,
                        colLabels=col_labels,
                        loc='center',
                        cellLoc='center',
                        colWidths=col_widths
                    )
                    table_obj.auto_set_font_size(False)
                    table_obj.set_fontsize(8)
                    table_obj.scale(1, 1.5)

                    for (row_idx, col_idx), cell in table_obj.get_celld().items():
                        if row_idx == 0:
                            cell.set_text_props(weight='bold', color='white')
                            cell.set_facecolor('#1a222d')
                        else:
                            cell.set_text_props(color='#111111')
                            if col_idx == 5:
                                status_val = cell_text[row_idx - 1][5]
                                if status_val == "Zaakceptowano":
                                    cell.set_facecolor('#d4edda')
                                    cell.set_text_props(color='#155724', weight='bold')
                                else:
                                    cell.set_facecolor('#f8d7da')
                                    cell.set_text_props(color='#721c24', weight='bold')

                    page_num = (i // rows_per_page) + 1
                    total_pages = (len(all_data) + rows_per_page - 1) // rows_per_page
                    fig_table.suptitle(f"Tabela Parametrów QC Wyników - Część {page_num} z {total_pages}",
                                       fontsize=12, weight='bold', y=0.96)

                    fig_table.tight_layout()
                    pdf.savefig(fig_table)

            QMessageBox.information(self, "Sukces", f"Pełny wielostronicowy raport PDF został zapisany w:\n{save_path}")
        except Exception as e:
            QMessageBox.critical(self, "Błąd", f"Nie udało się wygenerować pełnego raportu PDF:\n{e}")

    def export_fasta(self):
        save_path, _ = QFileDialog.getSaveFileName(self, "Zapisz czyste FASTA", "oczyszczone.fasta",
                                                   "Pliki FASTA (*.fasta *.fa)")
        if not save_path: return
        try:
            with open(save_path, "w", encoding="utf-8") as f:
                for r in self.accepted_records:
                    f.write(f">{r['id']}\n{r['sequence']}\n")
            QMessageBox.information(self, "Sukces", "Poprawne sekwencje zostały zapisane do pliku FASTA.")
        except Exception as e:
            QMessageBox.critical(self, "Błąd", f"Nie udało się zapisać pliku: {e}")


if __name__ == '__main__':
    app = QApplication(sys.argv)
    app.setStyle("Fusion")
    window = PhageQCApp()
    window.show()
    sys.exit(app.exec())